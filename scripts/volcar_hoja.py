# -*- coding: utf-8 -*-
"""Vuelca la plantilla .xlsx a JSON para poder probar el backend fuera de Google."""
import json, datetime
from openpyxl import load_workbook

wb = load_workbook('2casas-hoja-control.xlsx')
salida = {}
for nombre in wb.sheetnames:
    if nombre == 'LÉEME':
        continue
    filas = []
    for fila in wb[nombre].iter_rows(values_only=True):
        filas.append([
            '' if c is None else (c.isoformat() if isinstance(c, (datetime.datetime, datetime.date)) else c)
            for c in fila
        ])
    salida[nombre] = filas
with open('/tmp/hoja.json', 'w', encoding='utf-8') as f:
    json.dump(salida, f, ensure_ascii=False)
print('volcadas', len(salida), 'pestañas')
