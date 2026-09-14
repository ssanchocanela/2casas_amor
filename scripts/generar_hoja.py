# -*- coding: utf-8 -*-
"""Genera la plantilla de la hoja de control de 2 Casas + Amor."""

import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

FUENTE = "Arial"

VERDE = "1B5E20"
GRIS = "455A64"
AZUL = "0D47A1"
AMARILLO = "FFF9C4"
NIÑO1 = "00838F"
NIÑA2 = "C2185B"

wb = Workbook()
wb.remove(wb.active)

borde = Border(bottom=Side(style="thin", color="BDBDBD"))


def hoja(nombre, cabeceras, anchos, color=GRIS, nota=None):
    ws = wb.create_sheet(nombre)
    fila = 1
    if nota:
        ws.cell(row=1, column=1, value=nota).font = Font(
            name=FUENTE, size=9, italic=True, color="616161"
        )
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(cabeceras), 3))
        ws.row_dimensions[1].height = 24
        fila = 2
    for i, c in enumerate(cabeceras, start=1):
        celda = ws.cell(row=fila, column=i, value=c)
        celda.font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=color)
        celda.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[fila].height = 22
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)
    ws.auto_filter.ref = (
        f"A{fila}:{get_column_letter(len(cabeceras))}{fila}"
    )
    return ws, fila


def escribir(ws, fila_cab, filas, ejemplo_desde=None):
    for j, datos in enumerate(filas):
        r = fila_cab + 1 + j
        es_ejemplo = ejemplo_desde is not None and j >= ejemplo_desde
        for i, v in enumerate(datos, start=1):
            celda = ws.cell(row=r, column=i, value=v)
            celda.font = Font(name=FUENTE, size=10)
            celda.border = borde
            celda.alignment = Alignment(vertical="center", wrap_text=False)
            if es_ejemplo:
                celda.fill = PatternFill("solid", fgColor=AMARILLO)


# ─────────────────────────────────────────── LÉEME
ws = wb.create_sheet("LÉEME")
ws.column_dimensions["A"].width = 118
texto = [
    ("2 CASAS + AMOR — Hoja de control", 16, True, VERDE),
    ("", 10, False, "000000"),
    ("Esta hoja es la base de datos de la aplicación. Por eso parece aburrida: es lo que", 10, False, "000000"),
    ("hace que la app pueda leerla sin romperse, y lo que permitirá pasarla a una base de", 10, False, "000000"),
    ("datos de verdad el día que haga falta.", 10, False, "000000"),
    ("", 10, False, "000000"),
    ("Se empieza de cero con el curso 2026-2027. El histórico de la hoja anterior no se", 10, False, "000000"),
    ("vuelca: esa hoja se deja en solo lectura, renombrada con 'ARCHIVO —' delante.", 10, False, "000000"),
    ("", 10, False, "000000"),
    ("QUÉ SE RELLENA A MANO, UNA VEZ", 12, True, AZUL),
    ("   config · personas · ninos · custodia_mapeo · horario_centro · horario_lectivo ·", 10, False, "000000"),
    ("   extraescolares · calendario_escolar · listas", 10, False, "000000"),
    ("   Ya vienen rellenas con vuestros datos. Revisad y completad lo que falte.", 10, False, "000000"),
    ("", 10, False, "000000"),
    ("QUÉ SE RELLENA EN EL DÍA A DÍA", 12, True, AZUL),
    ("   menu_casa · registro_nino · salud_eventos · tareas · lecturas · eventos · menu_comedor", 10, False, "000000"),
    ("   Desde la app, o a mano aquí si preferís.", 10, False, "000000"),
    ("", 10, False, "000000"),
    ("QUÉ NO SE TOCA", 12, True, AZUL),
    ("   log — lo escribe la app sola. Es el registro de quién cambió qué y cuándo.", 10, False, "000000"),
    ("", 10, False, "000000"),
    ("LAS CINCO REGLAS", 12, True, VERDE),
    ("   1. Fila 1 (o 2 si hay nota) son las cabeceras. No se renombran ni se reordenan nunca.", 10, False, "000000"),
    ("   2. Un registro por fila. Ni celdas combinadas, ni filas de título, ni columnas en blanco.", 10, False, "000000"),
    ("   3. Nada se borra: se rellena 'eliminado_en' con la fecha. El historial se queda.", 10, False, "000000"),
    ("   4. Fechas AAAA-MM-DD y horas HH:MM, siempre. Da igual en qué móvil se abra.", 10, False, "000000"),
    ("   5. Para añadir una columna, va al final y se sube 'version_esquema' en config.", 10, False, "000000"),
    ("", 10, False, "000000"),
    ("LAS FILAS AMARILLAS SON EJEMPLOS", 12, True, "F57F17"),
    ("   Enseñan el formato esperado. Borradlas cuando empecéis a usar la hoja de verdad.", 10, False, "000000"),
    ("", 10, False, "000000"),
    ("IDENTIFICADORES", 12, True, AZUL),
    ("   prefijo_AAAAMMDD_xxxx  →  reg_20260913_a7f3. Los genera la app. A mano, vale", 10, False, "000000"),
    ("   cualquier cosa que no exista ya.", 10, False, "000000"),
    ("", 10, False, "000000"),
    ("POR QUÉ menu_casa Y registro_nino ESTÁN SEPARADOS", 12, True, VERDE),
    ("   Porque son dos cosas distintas. La cena se cocina una vez para los dos: es de la casa.", 10, False, "000000"),
    ("   La caca, la ducha o cómo ha comido cada uno son de cada niño. Meterlo todo en una", 10, False, "000000"),
    ("   misma fila obliga a inventar una columna nueva por cada niño y por cada cosa, que es", 10, False, "000000"),
    ("   justo lo que pasaba en la hoja anterior con 'Caca Niña2' y 'Lavado pelo Niño1'.", 10, False, "000000"),
]
for i, (t, s, b, c) in enumerate(texto, start=1):
    celda = ws.cell(row=i, column=1, value=t)
    celda.font = Font(name=FUENTE, size=s, bold=b, color=c)

# ─────────────────────────────────────────── config
ws, f = hoja(
    "config", ["clave", "valor", "notas"], [26, 46, 60], VERDE,
    "Parámetros generales. Una fila por parámetro.",
)
escribir(ws, f, [
    ["version_esquema", "1", "La app avisa si no coincide con la que espera"],
    ["curso_escolar", "2026-2027", ""],
    ["zona_horaria", "Europe/Madrid", ""],
    ["centro_escolar", "Colegio Ejemplo", "COMPLETAR: nombre del colegio"],
    ["calendar_id_custodia", "", "COMPLETAR: id del calendario de Google de la custodia"],
    ["calendar_id_colegio", "", "Opcional, si tenéis uno aparte"],
    ["carpeta_menus", "", "Opcional: id de la carpeta de Drive con el menú del comedor, una imagen por mes (AAAA-MM.jpg)"],
    ["entorno", "produccion", "En la copia de pruebas, poner: pruebas"],
])

# ─────────────────────────────────────────── personas
ws, f = hoja(
    "personas",
    ["id", "nombre", "rol", "email", "color_hex", "activo"],
    [16, 18, 14, 32, 12, 10], VERDE,
    "Quién puede aparecer como responsable de una casa o como autor de un apunte.",
)
escribir(ws, f, [
    ["per_papa", "Papá", "progenitor", "", "#2E7D32", "si"],
    ["per_mama", "Mamá", "progenitor", "", "#6A1B9A", "si"],
    ["per_abu", "Abuela", "cuidador", "", "#EF6C00", "si"],
])

# ─────────────────────────────────────────── ninos
ws, f = hoja(
    "ninos",
    ["id", "nombre", "fecha_nacimiento", "curso", "grupo", "centro",
     "color_hex", "alergias", "notas_medicas", "activo"],
    [14, 14, 18, 16, 10, 22, 12, 26, 34, 10], VERDE,
    "COMPLETAR las fechas de nacimiento, alergias y notas médicas.",
)
escribir(ws, f, [
    ["nin_nino1", "Niño1", "", "1º de Primaria", "A", "", f"#{NIÑO1}", "", "", "si"],
    ["nin_nina2", "Niña2", "", "Infantil 4", "B", "", f"#{NIÑA2}", "", "", "si"],
])

# ─────────────────────────────────────────── custodia_mapeo
ws, f = hoja(
    "custodia_mapeo",
    ["texto_evento", "coincidencia", "persona_id", "nino_id", "notas"],
    [28, 16, 16, 16, 46], VERDE,
    "Cómo traducir los títulos de vuestro Google Calendar a quién tiene a los niños. "
    "COMPLETAR con los títulos que usáis de verdad.",
)
escribir(ws, f, [
    ["PAPA", "contiene", "per_papa", "",
     "EJEMPLO — texto_evento: el código del evento en el calendario; persona_id: el id de 'personas'"],
    ["MAMA", "contiene", "per_mama", "", "EJEMPLO — nino_id vacío significa: los dos"],
], ejemplo_desde=0)

# ─────────────────────────────────────────── horario_centro
ws, f = hoja(
    "horario_centro",
    ["id", "nino_id", "curso_escolar", "entrada_desde", "entrada_hasta",
     "salida_desde", "salida_hasta", "puerta", "notas", "activo"],
    [16, 14, 15, 15, 15, 14, 14, 14, 42, 10], AZUL,
    "Horario de entrada y salida del colegio. La app calcula la hora real de recogida "
    "sumando las extraescolares de cada día.",
)
escribir(ws, f, [
    ["hce_nino1", "nin_nino1", "2026-2027", "09:00", "09:00", "17:00", "17:00",
     "Principal", "", "si"],
    ["hce_nina2", "nin_nina2", "2026-2027", "09:00", "09:00", "17:00", "17:00",
     "Principal", "", "si"],
])

# ─────────────────────────────────────────── horario_lectivo
ws, f = hoja(
    "horario_lectivo",
    ["id", "nino_id", "curso_escolar", "dia_semana", "orden", "hora_inicio",
     "hora_fin", "asignatura", "equipamiento", "aula", "profesor", "activo"],
    [18, 14, 15, 13, 8, 13, 12, 22, 15, 10, 16, 10], AZUL,
    "Horario de clases. 1=lunes … 5=viernes. 'equipamiento' es lo que hay que llevar ese día: "
    "CH chándal, PS piscina. Sale de scripts/datos/horario_lectivo_2026-2027.json.",
)
# El horario de verdad vive en scripts/datos/: es la misma fuente que usa
# scripts/cargar_horario.js para volcarlo en una hoja ya en marcha.
with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "datos",
                       "horario_lectivo_2026-2027.json"), encoding="utf-8") as fh:
    HORARIO = json.load(fh)
filas_horario = []
for nino_id, nino in HORARIO["ninos"].items():
    for dia, clases in nino["dias"].items():
        for i, clase in enumerate(clases):
            hora, asignatura = clase[0], clase[1]
            equipamiento = clase[2] if len(clase) > 2 else ""
            fin_clase = clases[i + 1][0] if i + 1 < len(clases) else nino["salida"]
            filas_horario.append([
                f"hor_{nino_id.replace('nin_', '')}_{dia}_{hora.replace(':', '')}", nino_id,
                HORARIO["curso_escolar"], int(dia), i + 1, hora, fin_clase, asignatura,
                equipamiento, "", "", "si",
            ])
escribir(ws, f, filas_horario)

# ─────────────────────────────────────────── extraescolares
ws, f = hoja(
    "extraescolares",
    ["id", "nino_id", "curso_escolar", "dia_semana", "actividad", "hora_inicio",
     "hora_fin", "equipamiento", "lugar", "fecha_desde", "fecha_hasta",
     "coste_mes", "pagado_por", "notas", "activo"],
    [20, 14, 15, 13, 24, 13, 12, 15, 16, 14, 14, 12, 14, 40, 10], AZUL,
    "Actividades de después del colegio. La hora_fin es la hora real de recogida ese día. "
    "'fecha_desde' es a partir de cuándo empieza: la primera sesión es el primer día de la "
    "semana que corresponda en esa fecha o después. Todas arrancan pasado el 01/10.",
)
escribir(ws, f, [
    ["ext_nino1_futbol_ma", "nin_nino1", "2026-2027", 2, "Fútbol", "", "18:00",
     "", "", "2026-10-01", "", "", "", "En el papel no figuraba fecha; se asume el inicio general de octubre", "si"],
    ["ext_nino1_futbol_ju", "nin_nino1", "2026-2027", 4, "Fútbol", "", "18:00",
     "", "", "2026-10-02", "", "", "", "Primer jueves a partir del 02/10: el 08/10", "si"],
    ["ext_nino1_robotica", "nin_nino1", "2026-2027", 5, "Robótica", "14:00", "15:00",
     "", "", "2026-10-01", "", "", "", "Es en horario de mediodía, no a la salida", "si"],
    ["ext_nina2_gimnasia", "nin_nina2", "2026-2027", 2, "Gimnasia", "", "18:15",
     "", "", "2026-10-01", "", "", "", "Primer martes a partir del 01/10: el 06/10", "si"],
    ["ext_nina2_baile", "nin_nina2", "2026-2027", 3, "Baile", "", "18:15",
     "", "", "2026-10-07", "", "", "", "", "si"],
])

# ─────────────────────────────────────────── calendario_escolar
ws, f = hoja(
    "calendario_escolar",
    ["id", "fecha", "ambito", "tipo", "descripcion"],
    [20, 14, 16, 18, 52], AZUL,
    "Festivos, días no lectivos, vacaciones y jornadas reducidas. "
    "COMPLETAR con el calendario que os dé el colegio.",
)
escribir(ws, f, [
    ["cal_20261012", "2026-10-12", "centro", "festivo", "EJEMPLO — Fiesta Nacional"],
], ejemplo_desde=0)

# ─────────────────────────────────────────── menu_casa
ws, f = hoja(
    "menu_casa",
    ["id", "creado_en", "fecha", "casa_persona_id", "comida", "cena", "fruta",
     "notas", "autor_id", "eliminado_en"],
    [20, 18, 13, 17, 36, 36, 18, 40, 14, 14], VERDE,
    "Qué se ha cocinado ese día en esa casa. Una fila por día y casa: la comida es de la "
    "casa, no de cada niño. Cómo comió cada uno va en registro_nino.",
)
escribir(ws, f, [
    ["men_20260915_x1", "2026-09-15 21:40", "2026-09-15", "per_papa",
     "EJEMPLO — Gambas, osobuco y espaguetis", "Dorada con cebolla y patatas", "Naranja",
     "Casi no cenan nada, se quedan fritos", "per_mama", ""],
], ejemplo_desde=0)

# ─────────────────────────────────────────── registro_nino
ws, f = hoja(
    "registro_nino",
    ["id", "creado_en", "fecha", "hora", "nino_id", "tipo", "valor", "escala",
     "casa_persona_id", "notas", "autor_id", "eliminado_en"],
    [20, 18, 13, 8, 14, 16, 30, 9, 17, 34, 14, 14], VERDE,
    "Una fila por cada cosa que le pasa a UN niño. 'escala' va de 1 a 5 en las comidas "
    "(1 nada, 5 todo) y en el ánimo. En caca o ducha, 'valor' es si / no — apuntar el 'no' "
    "también es información.",
)
escribir(ws, f, [
    ["reg_20260915_a1", "2026-09-15 21:41", "2026-09-15", "", "nin_nina2", "caca", "no", None,
     "per_papa", "EJEMPLO — el 'no' también se apunta", "per_mama", ""],
    ["reg_20260915_a2", "2026-09-15 21:41", "2026-09-15", "", "nin_nina2", "ducha", "si", None,
     "per_papa", "", "per_mama", ""],
    ["reg_20260915_a3", "2026-09-15 21:41", "2026-09-15", "", "nin_nina2", "lavado_pelo", "si", None,
     "per_papa", "", "per_mama", ""],
    ["reg_20260915_a4", "2026-09-15 21:42", "2026-09-15", "", "nin_nino1", "cena", "", 2,
     "per_papa", "Se quedó frito en la mesa", "per_mama", ""],
], ejemplo_desde=0)

# ─────────────────────────────────────────── salud_eventos
ws, f = hoja(
    "salud_eventos",
    ["id", "creado_en", "fecha", "nino_id", "tipo", "titulo", "descripcion",
     "profesional", "centro", "diagnostico", "tratamiento", "desde", "hasta",
     "adjunto_url", "autor_id", "eliminado_en"],
    [20, 18, 13, 14, 14, 28, 40, 22, 22, 28, 40, 13, 13, 26, 14, 14], "AD1457",
    "El control médico. 'desde' y 'hasta' marcan la pauta de un tratamiento: "
    "la app la recuerda mientras esté activa.",
)
escribir(ws, f, [
    ["sal_20260910_b1", "2026-09-10 18:00", "2026-09-10", "nin_nina2", "consulta",
     "EJEMPLO — Revisión pediatra", "", "", "", "", "", "", "", "", "per_mama", ""],
], ejemplo_desde=0)

# ─────────────────────────────────────────── salud_medidas
ws, f = hoja(
    "salud_medidas",
    ["id", "fecha", "nino_id", "peso_kg", "talla_cm", "notas", "autor_id"],
    [20, 13, 14, 11, 11, 40, 14], "AD1457",
    "Peso y talla, para poder ver la evolución.",
)
escribir(ws, f, [
    ["med_20260910_c1", "2026-09-10", "nin_nina2", 17.4, 104, "EJEMPLO", "per_mama"],
], ejemplo_desde=0)

# ─────────────────────────────────────────── tareas
ws, f = hoja(
    "tareas",
    ["id", "creado_en", "nino_id", "asignatura", "titulo", "descripcion", "tipo",
     "fecha_asignada", "fecha_entrega", "estado", "prioridad", "completada_en",
     "autor_id", "eliminado_en"],
    [20, 18, 14, 20, 32, 40, 14, 15, 15, 13, 12, 18, 14, 14], AZUL,
    "Deberes, exámenes, trabajos y entregas. 'fecha_entrega' es lo que se pinta en el calendario.",
)
escribir(ws, f, [
    ["tar_20260914_d1", "2026-09-14 19:00", "nin_nino1", "Medi", "EJEMPLO — Mural del cuerpo humano",
     "", "trabajo", "2026-09-14", "2026-09-21", "pendiente", "normal", "", "per_papa", ""],
], ejemplo_desde=0)

# ─────────────────────────────────────────── lecturas
ws, f = hoja(
    "lecturas",
    ["id", "nino_id", "titulo", "autor", "fecha_inicio", "fecha_fin",
     "paginas_totales", "paginas_leidas", "minutos", "valoracion", "comentario",
     "autor_id", "eliminado_en"],
    [20, 14, 34, 24, 14, 13, 15, 15, 10, 12, 40, 14, 14], AZUL,
    "Qué está leyendo cada uno. 'valoracion' de 1 a 5.",
)
escribir(ws, f, [
    ["lec_20260913_e1", "nin_nino1", "EJEMPLO — El monstruo de colores", "Anna Llenas",
     "2026-09-13", "", 48, 12, 15, None, "", "per_papa", ""],
], ejemplo_desde=0)

# ─────────────────────────────────────────── eventos
ws, f = hoja(
    "eventos",
    ["id", "nino_id", "titulo", "tipo", "fecha_inicio", "fecha_fin", "hora_inicio",
     "hora_fin", "lugar", "coste", "pagado_por", "notas", "calendar_event_id",
     "autor_id", "eliminado_en"],
    [20, 14, 32, 16, 14, 13, 13, 12, 24, 10, 14, 44, 26, 14, 14], AZUL,
    "Colonias, excursiones, reuniones con el tutor, cumpleaños. "
    "nino_id vacío = afecta a los dos. En 'notas', qué hay que llevar.",
)
escribir(ws, f, [
    ["eve_20261120_f1", "nin_nino1", "EJEMPLO — Colonias 1º de Primaria", "colonias",
     "2026-11-20", "2026-11-22", "", "", "", None, "", "Mochila, saco y linterna", "",
     "per_papa", ""],
], ejemplo_desde=0)

# ─────────────────────────────────────────── menu_comedor
ws, f = hoja(
    "menu_comedor",
    ["id", "fecha", "centro", "primero", "segundo", "postre", "alergenos",
     "origen", "autor_id"],
    [20, 13, 22, 34, 34, 20, 24, 12, 14], VERDE,
    "El menú del comedor del colegio, mes a mes. Se vuelca desde la app cuando el colegio lo manda.",
)
escribir(ws, f, [
    ["mnu_20260914_g1", "2026-09-14", "", "EJEMPLO — Arròs a la cassola", "Truita de patates",
     "Fruita", "", "manual", "per_papa"],
], ejemplo_desde=0)

# ─────────────────────────────────────────── listas
ws, f = hoja(
    "listas",
    ["rol", "registro_tipo", "salud_tipo", "tarea_tipo", "tarea_estado",
     "evento_tipo", "calendario_tipo", "coincidencia", "menu_origen",
     "prioridad", "equipamiento", "si_no"],
    [14, 16, 14, 13, 14, 16, 18, 14, 13, 11, 14, 8], GRIS,
    "Los valores válidos de cada campo de lista. Cambiar algo aquí lo cambia en toda la hoja.",
)
listas = {
    "rol": ["progenitor", "cuidador", "nino"],
    "registro_tipo": ["desayuno", "comida", "merienda", "cena", "caca", "pipi",
                      "ducha", "bano", "lavado_pelo", "medicacion", "animo",
                      "sueno", "observacion"],
    "salud_tipo": ["consulta", "urgencia", "vacuna", "sintoma", "medicacion",
                   "alergia", "revision", "dentista", "otro"],
    "tarea_tipo": ["deber", "examen", "trabajo", "proyecto", "entrega", "lectura"],
    "tarea_estado": ["pendiente", "en_curso", "hecha", "entregada"],
    "evento_tipo": ["colonias", "excursion", "extraescolar", "reunion_tutor",
                    "cumple", "otro"],
    "calendario_tipo": ["festivo", "no_lectivo", "vacaciones", "jornada_reducida",
                        "evento"],
    "coincidencia": ["exacta", "contiene", "empieza_por"],
    "menu_origen": ["manual", "foto", "pdf"],
    "prioridad": ["normal", "alta"],
    "equipamiento": ["CH", "PS", "CH+PS"],
    "si_no": ["si", "no"],
}
cols = list(listas.keys())
maxlen = max(len(v) for v in listas.values())
for ci, nombre in enumerate(cols, start=1):
    for ri, valor in enumerate(listas[nombre]):
        celda = ws.cell(row=f + 1 + ri, column=ci, value=valor)
        celda.font = Font(name=FUENTE, size=10)

# nota al pie de listas
ws.cell(row=f + maxlen + 3, column=1,
        value="CH = chándal · PS = material de piscina · CH+PS = las dos cosas"
        ).font = Font(name=FUENTE, size=9, italic=True, color="616161")

FILA_LISTAS = f + 1

# ─────────────────────────────────────────── log
ws, f = hoja(
    "log",
    ["id", "timestamp", "autor_id", "accion", "pestana", "registro_id", "campo",
     "valor_anterior", "valor_nuevo", "origen"],
    [20, 18, 14, 11, 18, 20, 18, 28, 28, 11], "424242",
    "NO TOCAR. Lo escribe la app sola: quién cambió qué y cuándo.",
)

# ─────────────────────────────────────────── validaciones
def validar(nombre_hoja, columna, lista, primera_fila, ultima_fila=2000):
    col_lista = get_column_letter(cols.index(lista) + 1)
    n = len(listas[lista])
    dv = DataValidation(
        type="list",
        formula1=f"=listas!${col_lista}${FILA_LISTAS}:${col_lista}${FILA_LISTAS + n - 1}",
        allow_blank=True,
        showDropDown=False,
    )
    dv.error = "Ese valor no está en la pestaña listas."
    dv.errorTitle = "Valor no válido"
    h = wb[nombre_hoja]
    h.add_data_validation(dv)
    dv.add(f"{columna}{primera_fila}:{columna}{ultima_fila}")


validar("personas", "C", "rol", 3)
validar("custodia_mapeo", "B", "coincidencia", 3)
validar("horario_lectivo", "I", "equipamiento", 3)
validar("extraescolares", "H", "equipamiento", 3)
validar("calendario_escolar", "D", "calendario_tipo", 3)
validar("registro_nino", "F", "registro_tipo", 3)
validar("salud_eventos", "E", "salud_tipo", 3)
validar("tareas", "G", "tarea_tipo", 3)
validar("tareas", "J", "tarea_estado", 3)
validar("tareas", "K", "prioridad", 3)
validar("eventos", "D", "evento_tipo", 3)
validar("menu_comedor", "H", "menu_origen", 3)

# LÉEME primero
wb.move_sheet("LÉEME", offset=-(len(wb.sheetnames) - 1))

wb.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "2casas-hoja-control.xlsx"))
print("OK ·", len(wb.sheetnames), "pestañas:", ", ".join(wb.sheetnames))
